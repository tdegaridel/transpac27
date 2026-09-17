#!/usr/bin/env python3
"""
Régénère leg1.csv et leg2.csv à partir du classeur maître 20260916_TransPac_route.xlsx.

Le classeur (.xlsx) est la source de vérité : il porte les formules.
Les CSV en sont dérivés, pour un suivi des modifications lisible ligne à ligne sous Git.

Usage :
    python export_csv.py

Prérequis :
    pip install openpyxl

Note : le classeur doit avoir été recalculé (ouvert/enregistré dans un tableur, ou
via LibreOffice) pour que les cellules à formule contiennent leur valeur en cache.
Sans recalcul, les colonnes calculées (Ops total) peuvent apparaître vides.
"""

import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl manquant. Installe-le avec : pip install openpyxl")

XLSX = "20260916_TransPac_route.xlsx"

# Onglet du classeur -> fichier CSV de sortie
SHEETS = {
    "20260916_TransPac_leg1": "leg1.csv",
    "20260916_leg2": "leg2.csv",
}


def export(xlsx_path: Path) -> None:
    if not xlsx_path.exists():
        sys.exit(f"Fichier introuvable : {xlsx_path}")

    # data_only=True : on lit les valeurs en cache, pas les formules
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    for sheet_name, out_name in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"  ! onglet absent, ignoré : {sheet_name}")
            continue
        ws = wb[sheet_name]
        out_path = xlsx_path.parent / out_name
        with open(out_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if c is None else c for c in row])
        print(f"  \u2713 {out_name}")


if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    print("Export des CSV depuis", XLSX)
    export(base / XLSX)
    print("Terminé. Pense à : git add -A && git commit -m \"maj route\" && git push")
