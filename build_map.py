#!/usr/bin/env python3
"""
build_map.py — régénère TOUT à partir du classeur maître.

Source de vérité : 20260916_TransPac_route.xlsx
Sorties régénérées :
  - leg1.csv, leg2.csv                (export texte pour les diffs Git)
  - TransPac_carte_bathy_v2.html      (données réinjectées dans le bloc AUTOGEN)

La carte reste un fichier HTML autonome (ouvrable d'un double-clic, sans serveur) :
ce script réécrit seulement le tableau de stations entre les marqueurs
  // >>> AUTOGEN:DATA   ...   // <<< AUTOGEN:DATA
Le reste du HTML (style, logique Leaflet, noms de ZEE) n'est jamais touché.

Usage :
    python build_map.py

Prérequis :
    pip install openpyxl
    Le classeur doit avoir été recalculé (ouvert/enregistré dans un tableur)
    pour que la colonne « Ops total » soit renseignée dans les CSV.
"""

import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl manquant. Installe-le avec : pip install openpyxl")

XLSX = "20260916_TransPac_route.xlsx"
HTML = "TransPac_carte_bathy_v2.html"

SHEETS = {
    "20260916_TransPac_leg1": "leg1.csv",
    "20260916_leg2": "leg2.csv",
}

# Index de colonnes (1-based) dans le classeur -> voir en-têtes
COL = {
    "site_id": 3, "site": 4, "lon": 6, "lat": 7, "depth": 9,
    "CS": 11, "C1": 12, "C2": 13, "MN": 14, "NP": 15, "NM": 16,
    "CA": 17, "MT": 18, "type": 5, "note": 20,
}
INSTRUMENTS = ["CS", "C1", "C2", "MN", "NP", "NM", "CA", "MT"]


def export_csv(wb, base: Path) -> None:
    for sheet, out in SHEETS.items():
        ws = wb[sheet]
        with open(base / out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if c is None else c for c in row])
        print(f"  \u2713 {out}")


def js_val(x):
    """Formate une valeur Python en littéral JS."""
    if x is None or x == "":
        return "null"
    if isinstance(x, bool):
        return "true" if x else "false"
    if isinstance(x, (int, float)):
        return repr(x)
    s = str(x).replace("\\", "\\\\").replace('"', '\\"')
    return f'"{s}"'


def row_to_map_entry(ws, r):
    """Convertit une ligne du classeur en entrée carte [code,site,lon,lat,type,instr,depth,isNew]."""
    def g(key):
        return ws.cell(row=r, column=COL[key]).value

    sid = g("site_id")
    if sid in (None, ""):
        return None
    site = g("site") or ""
    lon = g("lon")
    lat = g("lat")
    if lon in (None, "") or lat in (None, ""):
        return None
    typ = g("type")
    is_port = (typ == "Port call")
    # instruments présents = colonnes horaires numériques > 0
    instr = []
    for code in INSTRUMENTS:
        v = g(code)
        if isinstance(v, (int, float)) and v > 0:
            instr.append(code)
    depth = g("depth")
    if isinstance(depth, (int, float)):
        depth = int(depth) if float(depth).is_integer() else depth
    note = str(g("note") or "")
    is_new = "AJOUTÉE" in note or "JOKER" in note
    return [
        str(sid), str(site), float(lon), float(lat),
        "port" if is_port else "op",
        "" if is_port else ",".join(instr),
        None if is_port else depth,
        bool(is_new),
    ]


def build_js_array(name, entries):
    lines = [f"const {name} = ["]
    for e in entries:
        lines.append(" [" + ",".join(js_val(v) for v in e) + "],")
    lines.append("];")
    return "\n".join(lines)


def update_html(wb, base: Path) -> None:
    leg1 = [row_to_map_entry(wb["20260916_TransPac_leg1"], r)
            for r in range(2, wb["20260916_TransPac_leg1"].max_row + 1)]
    leg2 = [row_to_map_entry(wb["20260916_leg2"], r)
            for r in range(2, wb["20260916_leg2"].max_row + 1)]
    leg1 = [e for e in leg1 if e]
    leg2 = [e for e in leg2 if e]

    block = (build_js_array("LEG1", leg1) + "\n" + build_js_array("LEG2", leg2))

    html_path = base / HTML
    html = html_path.read_text(encoding="utf-8")
    pattern = re.compile(
        r"(// >>> AUTOGEN:DATA\n).*?(\n// <<< AUTOGEN:DATA)",
        re.DOTALL,
    )
    if not pattern.search(html):
        sys.exit("Marqueurs AUTOGEN introuvables dans le HTML — abandon (aucune modif).")
    new_html = pattern.sub(lambda m: m.group(1) + block + m.group(2), html)
    html_path.write_text(new_html, encoding="utf-8")
    print(f"  \u2713 {HTML} ({len(leg1)} stations Leg 1, {len(leg2)} Leg 2)")


if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    xlsx_path = base / XLSX
    if not xlsx_path.exists():
        sys.exit(f"Fichier introuvable : {xlsx_path}")
    print("Régénération depuis", XLSX)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    export_csv(wb, base)
    update_html(wb, base)
    print('Terminé. git add -A && git commit -m "maj route" && git push')
